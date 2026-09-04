from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle
)


# ============================================================
# HELPER
# ============================================================

def clean_filename(text):
    """
    Convert a title into a safe filename.
    """

    if not text:
        return "polar_research_article"

    allowed = []

    for char in text:

        if char.isalnum():
            allowed.append(char)

        elif char in (" ", "-", "_"):
            allowed.append("_")

    filename = "".join(allowed)

    return filename[:80].strip("_")


# ============================================================
# DOCX BORDER
# ============================================================

def add_cell_border(cell, color="B7C9D6", size="8"):
    """
    Add a border around a table cell.
    """

    tc = cell._tc

    tcPr = tc.get_or_add_tcPr()

    tcBorders = tcPr.first_child_found_in("w:tcBorders")

    if tcBorders is None:

        tcBorders = OxmlElement("w:tcBorders")

        tcPr.append(tcBorders)

    for edge in ("top", "left", "bottom", "right"):

        tag = "w:" + edge

        element = tcBorders.find(qn(tag))

        if element is None:

            element = OxmlElement(tag)

            tcBorders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


# ============================================================
# DOCX GENERATOR
# ============================================================

def create_docx(
    title,
    content,
    query="",
    sources=None
):

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)


    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    title_paragraph = document.add_paragraph()

    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_paragraph.add_run(
        title
    )

    title_run.bold = True
    title_run.font.size = Pt(25)
    title_run.font.color.rgb = RGBColor(
        25,
        65,
        95
    )


    # --------------------------------------------------------
    # SUBTITLE
    # --------------------------------------------------------

    subtitle = document.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle_run = subtitle.add_run(
        "Polar Research Outreach Article"
    )

    subtitle_run.italic = True
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor(
        90,
        105,
        120
    )


    # --------------------------------------------------------
    # LINE
    # --------------------------------------------------------

    line = document.add_paragraph()

    run = line.add_run(
        "____________________________________________________________"
    )

    run.font.color.rgb = RGBColor(
        180,
        200,
        215
    )


    # --------------------------------------------------------
    # TOPIC
    # --------------------------------------------------------

    if query:

        topic = document.add_paragraph()

        topic_run = topic.add_run(
            "Research Topic: "
        )

        topic_run.bold = True

        topic.add_run(
            query
        )


    # --------------------------------------------------------
    # CONTENT
    # --------------------------------------------------------

    paragraphs = content.split("\n")

    for paragraph_text in paragraphs:

        text = paragraph_text.strip()

        if not text:
            continue


        # Heading detection

        if (
            text.startswith("# ")
            or text.startswith("## ")
            or text.isupper()
        ):

            heading_text = (
                text.replace("# ", "")
                .replace("## ", "")
            )

            heading = document.add_paragraph()

            heading_run = heading.add_run(
                heading_text
            )

            heading_run.bold = True
            heading_run.font.size = Pt(16)
            heading_run.font.color.rgb = RGBColor(
                25,
                65,
                95
            )

        else:

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(
                8
            )

            run = paragraph.add_run(
                text
            )

            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(
                35,
                45,
                55
            )


    # --------------------------------------------------------
    # SOURCES
    # --------------------------------------------------------

    if sources:

        document.add_paragraph()

        heading = document.add_paragraph()

        heading_run = heading.add_run(
            "Research Sources"
        )

        heading_run.bold = True
        heading_run.font.size = Pt(16)
        heading_run.font.color.rgb = RGBColor(
            25,
            65,
            95
        )


        table = document.add_table(
            rows=1,
            cols=2
        )

        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        table.style = "Table Grid"


        header = table.rows[0].cells

        header[0].text = "Source"
        header[1].text = "Research Record"


        for cell in header:

            add_cell_border(cell)

            for paragraph in cell.paragraphs:

                for run in paragraph.runs:

                    run.bold = True


        for source in sources:

            row = table.add_row().cells

            row[0].text = str(
                source.get(
                    "source",
                    "Research Archive"
                )
            )

            row[1].text = str(
                source.get(
                    "title",
                    "Research Record"
                )
            )

            for cell in row:

                add_cell_border(cell)

                cell.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                )


    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "Polar Research Knowledge Portal"
    )

    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(
        110,
        120,
        130
    )


    # --------------------------------------------------------
    # SAVE TO MEMORY
    # --------------------------------------------------------

    output = BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output


# ============================================================
# PDF GENERATOR
# ============================================================

def create_pdf(
    title,
    content,
    query="",
    sources=None
):

    output = BytesIO()


    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )


    styles = getSampleStyleSheet()


    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        fontSize=22,
        leading=27,
        alignment=TA_CENTER,
        textColor=colors.HexColor(
            "#19415F"
        ),
        spaceAfter=10
    )


    subtitle_style = ParagraphStyle(
        "Subtitle",
        parent=styles["Normal"],
        fontSize=10,
        alignment=TA_CENTER,
        textColor=colors.HexColor(
            "#687887"
        ),
        spaceAfter=20
    )


    heading_style = ParagraphStyle(
        "Heading",
        parent=styles["Heading2"],
        fontSize=15,
        leading=19,
        textColor=colors.HexColor(
            "#19415F"
        ),
        spaceBefore=14,
        spaceAfter=7
    )


    body_style = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontSize=10.5,
        leading=16,
        textColor=colors.HexColor(
            "#232D37"
        ),
        spaceAfter=8
    )


    story = []


    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    story.append(
        Paragraph(
            title,
            title_style
        )
    )


    story.append(
        Paragraph(
            "Polar Research Outreach Article",
            subtitle_style
        )
    )


    if query:

        story.append(
            Paragraph(
                "<b>Research Topic:</b> "
                + query,
                body_style
            )
        )


    story.append(
        Spacer(
            1,
            12
        )
    )


    # --------------------------------------------------------
    # CONTENT
    # --------------------------------------------------------

    for line in content.split("\n"):

        text = line.strip()

        if not text:
            continue


        if (
            text.startswith("# ")
            or text.startswith("## ")
            or text.isupper()
        ):

            heading = (
                text
                .replace("# ", "")
                .replace("## ", "")
            )

            story.append(
                Paragraph(
                    heading,
                    heading_style
                )
            )

        else:

            safe_text = (
                text
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            story.append(
                Paragraph(
                    safe_text,
                    body_style
                )
            )


    # --------------------------------------------------------
    # SOURCES
    # --------------------------------------------------------

    if sources:

        story.append(
            Paragraph(
                "Research Sources",
                heading_style
            )
        )


        table_data = [
            [
                "Source",
                "Research Record"
            ]
        ]


        for source in sources:

            table_data.append(
                [
                    str(
                        source.get(
                            "source",
                            "Research Archive"
                        )
                    ),

                    str(
                        source.get(
                            "title",
                            "Research Record"
                        )
                    )
                ]
            )


        table = Table(
            table_data,
            colWidths=[
                130,
                340
            ]
        )


        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#EAF1F6"
                        )
                    ),

                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#19415F"
                        )
                    ),

                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold"
                    ),

                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.6,
                        colors.HexColor(
                            "#B7C9D6"
                        )
                    ),

                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP"
                    ),

                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        9
                    ),

                    (
                        "TEXTCOLOR",
                        (0, 1),
                        (-1, -1),
                        colors.HexColor(
                            "#303A44"
                        )
                    ),

                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        7
                    ),

                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        7
                    )
                ]
            )
        )


        story.append(
            table
        )


    document.build(
        story
    )


    output.seek(0)

    return output