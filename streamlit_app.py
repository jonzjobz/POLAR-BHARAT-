import streamlit as st
import pandas as pd
import io
import re
import html

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image as PDFImage
)
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle
)
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch

from search_agent import (
    search_records,
    analyze_records,
    generate_outreach,
    generate_article,
    generate_research_image,
    get_source,
    get_title,
    clean_text
)


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="POLAR BHARAT",
    page_icon=None,
    layout="wide",
    initial_sidebar_state="expanded"
)


# ============================================================
# GLOBAL CSS
# Styling only — no HTML content is used to build the UI.
# ============================================================

st.markdown(
    """
    <style>

    /* --------------------------------------------------------
       GLOBAL
    -------------------------------------------------------- */

    .stApp {
        background-color: #f5f7f9;
    }

    .block-container {
        max-width: 1450px;
        padding-top: 2rem;
        padding-bottom: 4rem;
        padding-left: 3rem;
        padding-right: 3rem;
    }

    /*
       Force normal readable text throughout the application.
    */

    .stApp,
    .stApp p,
    .stApp label,
    .stApp span,
    .stApp div {
        color: #172b3a;
    }

    /*
       Main headings
    */

    h1,
    h2,
    h3,
    h4 {
        color: #0b2942 !important;
        font-weight: 700 !important;
    }

    h1 {
        font-size: 2.3rem !important;
    }

    h2 {
        font-size: 1.65rem !important;
    }

    h3 {
        font-size: 1.25rem !important;
    }

    /*
       Horizontal rules
    */

    hr {
        border: none !important;
        border-top: 1px solid #d8e0e6 !important;
        margin-top: 1.5rem !important;
        margin-bottom: 1.5rem !important;
    }


    /* --------------------------------------------------------
       SIDEBAR
    -------------------------------------------------------- */

    section[data-testid="stSidebar"] {
        background-color: #ffffff;
        border-right: 1px solid #d7e0e6;
    }

    section[data-testid="stSidebar"] > div {
        padding-top: 1.5rem;
    }

    section[data-testid="stSidebar"],
    section[data-testid="stSidebar"] p,
    section[data-testid="stSidebar"] span,
    section[data-testid="stSidebar"] label {
        color: #173247 !important;
    }

    section[data-testid="stSidebar"] h1,
    section[data-testid="stSidebar"] h2,
    section[data-testid="stSidebar"] h3 {
        color: #0b2942 !important;
    }


    /* --------------------------------------------------------
       RADIO NAVIGATION
    -------------------------------------------------------- */

    section[data-testid="stSidebar"]
    div[role="radiogroup"] {
        gap: 5px;
    }

    section[data-testid="stSidebar"]
    div[role="radiogroup"] label {
        background-color: transparent !important;
        border-radius: 7px;
        padding: 8px 10px;
    }

    section[data-testid="stSidebar"]
    div[role="radiogroup"] label:hover {
        background-color: #eef4f7 !important;
    }

    section[data-testid="stSidebar"]
    div[role="radiogroup"] label p {
        color: #173247 !important;
        font-size: 14px !important;
    }


    /* --------------------------------------------------------
       TEXT INPUTS
    -------------------------------------------------------- */

    input,
    textarea {
        color: #172b3a !important;
        background-color: #ffffff !important;
        caret-color: #173f5c !important;
    }

    input::placeholder,
    textarea::placeholder {
        color: #7a8d9a !important;
        opacity: 1 !important;
    }


    /* --------------------------------------------------------
       SELECTBOX
    -------------------------------------------------------- */

    div[data-baseweb="select"] {
        background-color: #ffffff !important;
    }

    div[data-baseweb="select"] > div {
        background-color: #ffffff !important;
        color: #172b3a !important;
        border-color: #c8d4dc !important;
    }

    div[data-baseweb="select"] span {
        color: #172b3a !important;
    }

    ul[role="listbox"] {
        background-color: #ffffff !important;
        border: 1px solid #c8d4dc !important;
    }

    ul[role="listbox"] li {
        color: #172b3a !important;
        background-color: #ffffff !important;
    }

    ul[role="listbox"] li:hover {
        background-color: #edf4f7 !important;
        color: #0b2942 !important;
    }


    /* --------------------------------------------------------
       BUTTONS
    -------------------------------------------------------- */

    .stButton > button {
        background-color: #ffffff !important;
        color: #173f5c !important;
        border: 1px solid #aebfca !important;
        border-radius: 7px !important;
        min-height: 40px !important;
        font-weight: 600 !important;
    }

    .stButton > button:hover {
        background-color: #eef4f7 !important;
        border-color: #173f5c !important;
        color: #0b2942 !important;
    }

    .stButton > button[kind="primary"] {
        background-color: #0b3a59 !important;
        color: #ffffff !important;
        border-color: #0b3a59 !important;
    }

    .stButton > button[kind="primary"]:hover {
        background-color: #092f48 !important;
        color: #ffffff !important;
    }

    /*
       Download buttons
    */

    .stDownloadButton > button {
        background-color: #ffffff !important;
        color: #173f5c !important;
        border: 1px solid #aebfca !important;
        border-radius: 7px !important;
        font-weight: 600 !important;
    }

    .stDownloadButton > button:hover {
        background-color: #eef4f7 !important;
        border-color: #173f5c !important;
    }


    /* --------------------------------------------------------
       METRICS
    -------------------------------------------------------- */

    [data-testid="stMetric"] {
        background-color: #ffffff !important;
        border: 1px solid #d6e0e6 !important;
        border-radius: 9px !important;
        padding: 18px !important;
        box-shadow: 0 2px 7px rgba(11, 41, 66, 0.04);
    }

    [data-testid="stMetricLabel"] {
        color: #647987 !important;
    }

    [data-testid="stMetricValue"] {
        color: #0b2942 !important;
    }


    /* --------------------------------------------------------
       EXPANDERS
    -------------------------------------------------------- */

    [data-testid="stExpander"] {
        background-color: #ffffff !important;
        border: 1px solid #d6e0e6 !important;
        border-radius: 8px !important;
    }

    [data-testid="stExpander"] summary {
        color: #173247 !important;
    }

    [data-testid="stExpander"] summary p {
        color: #173247 !important;
        font-weight: 600 !important;
    }


    /* --------------------------------------------------------
       ALERTS
    -------------------------------------------------------- */

    [data-testid="stAlert"] {
        border-radius: 8px !important;
    }

    [data-testid="stAlert"] p,
    [data-testid="stAlert"] span {
        color: #173247 !important;
    }


    /* --------------------------------------------------------
       DATAFRAME
    -------------------------------------------------------- */

    [data-testid="stDataFrame"] {
        border: 1px solid #d6e0e6 !important;
        border-radius: 8px !important;
        overflow: hidden;
    }


    /* --------------------------------------------------------
       CHECKBOX
    -------------------------------------------------------- */

    [data-testid="stCheckbox"] label p {
        color: #173247 !important;
    }


    /* --------------------------------------------------------
       CAPTIONS
    -------------------------------------------------------- */

    .stCaption {
        color: #667b89 !important;
    }


    /* --------------------------------------------------------
       LINKS
    -------------------------------------------------------- */

    a {
        color: #075a78 !important;
    }


    /* --------------------------------------------------------
       CARD-LIKE CONTAINERS
    -------------------------------------------------------- */

    div[data-testid="stVerticalBlockBorderWrapper"] {
        background-color: #ffffff;
        border-color: #d6e0e6 !important;
        border-radius: 9px !important;
    }


    /* --------------------------------------------------------
       FOOTER
    -------------------------------------------------------- */

    .footer-note {
        color: #718391;
        font-size: 0.85rem;
        text-align: center;
        padding-top: 25px;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# DATA FILES
# ============================================================

DATA_FILES = [
    "arctic_projects.csv",
    "arctic_team_members.csv",
    "npdc_datasets.csv",
    "npdc_prototype.csv",
    "npdc_station_metadata.csv",
    "polar_images.csv",
    "polar_publications.csv",
    "polar_southern_ocean_reports.csv"
]


# ============================================================
# COLLECTION NAMES
# ============================================================

COLLECTION_NAMES = {

    "arctic_projects.csv":
        "Arctic Projects",

    "arctic_team_members.csv":
        "Arctic Team Members",

    "npdc_datasets.csv":
        "NPDC Scientific Datasets",

    "npdc_prototype.csv":
        "NPDC Prototype Data",

    "npdc_station_metadata.csv":
        "NPDC Station Metadata",

    "polar_images.csv":
        "Polar Images",

    "polar_publications.csv":
        "Polar Publications",

    "polar_southern_ocean_reports.csv":
        "Southern Ocean Reports"
}


# ============================================================
# COLLECTION DESCRIPTIONS
# ============================================================

COLLECTION_DESCRIPTIONS = {

    "arctic_projects.csv":
        "Research projects and scientific activities associated with Arctic research.",

    "arctic_team_members.csv":
        "Researcher and team-member information associated with Arctic projects.",

    "npdc_datasets.csv":
        "Scientific datasets available through the National Polar Data Centre.",

    "npdc_prototype.csv":
        "Prototype research records collected from the NPDC ecosystem.",

    "npdc_station_metadata.csv":
        "Metadata describing polar research stations and related information.",

    "polar_images.csv":
        "Polar research imagery and associated metadata.",

    "polar_publications.csv":
        "Research publications and associated scientific metadata.",

    "polar_southern_ocean_reports.csv":
        "Reports and research records related to Southern Ocean activities."
}


# ============================================================
# MAIN RESEARCH TOPICS
# ============================================================

MAIN_TOPICS = [
    "Arctic Research",
    "Antarctic Research",
    "Climate Change",
    "Atmospheric Science",
    "Oceanography",
    "Cryosphere",
    "Biodiversity",
    "Geology",
    "Polar Biology",
    "Remote Sensing",
    "Glaciology",
    "Southern Ocean"
]


# ============================================================
# LOAD DATA
# ============================================================

@st.cache_data
def load_all_data():

    data = {}

    for filename in DATA_FILES:

        try:

            df = pd.read_csv(
                filename,
                encoding="utf-8",
                on_bad_lines="skip"
            )

            data[filename] = df

        except Exception:

            data[filename] = pd.DataFrame()

    return data


data = load_all_data()


# ============================================================
# SOURCE IDENTIFICATION
# ============================================================

def source_from_filename(filename):

    filename = filename.lower()

    if "npdc" in filename:
        return "NPDC"

    if (
        "arctic" in filename
        or "publication" in filename
        or "southern" in filename
        or "image" in filename
    ):
        return "NCPOR"

    return "Primary Dataset"


# ============================================================
# SAFE VALUE
# ============================================================

def safe_value(row, column):

    if column not in row.index:
        return ""

    value = row.get(column, "")

    try:

        if pd.isna(value):
            return ""

    except Exception:
        pass

    return str(value).strip()


# ============================================================
# ARTICLE PARSER
# ============================================================

def parse_article(article_text):

    sections = {}

    current_section = None
    current_content = []

    section_names = [
        "TITLE",
        "SUBTITLE",
        "INTRODUCTION",
        "RESEARCH HIGHLIGHTS",
        "MAIN STORY",
        "GEOGRAPHIC CONTEXT",
        "WHY IT MATTERS",
        "CONCLUSION",
        "SOURCE RECORDS",
        "IMAGE DESCRIPTION"
    ]

    for line in article_text.splitlines():

        clean_line = line.strip()

        matched_section = None

        for section in section_names:

            if clean_line.upper().startswith(
                section + ":"
            ):

                matched_section = section
                break

        if matched_section:

            if current_section:

                sections[current_section] = (
                    "\n".join(
                        current_content
                    ).strip()
                )

            current_section = matched_section

            content_after_colon = clean_line[
                len(matched_section) + 1:
            ].strip()

            current_content = []

            if content_after_colon:

                current_content.append(
                    content_after_colon
                )

        else:

            if current_section:

                current_content.append(line)

    if current_section:

        sections[current_section] = (
            "\n".join(
                current_content
            ).strip()
        )

    return sections


# ============================================================
# DOCX GENERATOR
# ============================================================

def create_docx(
    article_text,
    image_bytes=None
):

    sections = parse_article(article_text)

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = sections.get(
        "TITLE",
        "Polar Research Article"
    )

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(title)

    run.bold = True
    run.font.size = Pt(24)

    subtitle = sections.get(
        "SUBTITLE",
        ""
    )

    if subtitle:

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(subtitle)

        run.italic = True
        run.font.size = Pt(12)

    if image_bytes:

        image_stream = io.BytesIO(image_bytes)

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()

        run.add_picture(
            image_stream,
            width=Inches(6.2)
        )

    display_sections = [
        "INTRODUCTION",
        "RESEARCH HIGHLIGHTS",
        "MAIN STORY",
        "GEOGRAPHIC CONTEXT",
        "WHY IT MATTERS",
        "CONCLUSION"
    ]

    for section_name in display_sections:

        content = sections.get(
            section_name,
            ""
        )

        if not content:
            continue

        heading = document.add_paragraph()

        run = heading.add_run(
            section_name.title()
        )

        run.bold = True
        run.font.size = Pt(15)

        for line in content.splitlines():

            line = line.strip()

            if not line:
                continue

            if line.startswith("-"):

                paragraph = document.add_paragraph(
                    style="List Bullet"
                )

                paragraph.add_run(
                    line.lstrip("-").strip()
                )

            elif line.startswith("•"):

                paragraph = document.add_paragraph(
                    style="List Bullet"
                )

                paragraph.add_run(
                    line.lstrip("•").strip()
                )

            else:

                document.add_paragraph(line)

    sources = sections.get(
        "SOURCE RECORDS",
        ""
    )

    if sources:

        heading = document.add_paragraph()

        run = heading.add_run(
            "Research Evidence & Sources"
        )

        run.bold = True
        run.font.size = Pt(15)

        for line in sources.splitlines():

            line = line.strip()

            if line:

                document.add_paragraph(
                    line,
                    style="List Bullet"
                )

    footer = section.footer.paragraphs[0]

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer.add_run(
        "POLAR BHARAT | Evidence-based generated content"
    )

    footer_run.font.size = Pt(8)

    output = io.BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# PDF GENERATOR
# ============================================================

def create_pdf(
    article_text,
    image_bytes=None
):

    sections = parse_article(article_text)

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ArticleTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=22,
        leading=26,
        spaceAfter=10
    )

    subtitle_style = ParagraphStyle(
        "ArticleSubtitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=11,
        leading=16,
        spaceAfter=18
    )

    heading_style = ParagraphStyle(
        "ArticleHeading",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        spaceBefore=14,
        spaceAfter=7
    )

    body_style = ParagraphStyle(
        "ArticleBody",
        parent=styles["BodyText"],
        fontSize=10.5,
        leading=16,
        spaceAfter=8
    )

    story = []

    title = sections.get(
        "TITLE",
        "Polar Research Article"
    )

    story.append(
        Paragraph(
            html.escape(title),
            title_style
        )
    )

    subtitle = sections.get(
        "SUBTITLE",
        ""
    )

    if subtitle:

        story.append(
            Paragraph(
                html.escape(subtitle),
                subtitle_style
            )
        )

    if image_bytes:

        image_stream = io.BytesIO(image_bytes)

        image = PDFImage(
            image_stream,
            width=6.5 * inch,
            height=3.65 * inch
        )

        story.append(image)

        story.append(
            Spacer(1, 15)
        )

    display_sections = [
        "INTRODUCTION",
        "RESEARCH HIGHLIGHTS",
        "MAIN STORY",
        "GEOGRAPHIC CONTEXT",
        "WHY IT MATTERS",
        "CONCLUSION"
    ]

    for section_name in display_sections:

        content = sections.get(
            section_name,
            ""
        )

        if not content:
            continue

        story.append(
            Paragraph(
                section_name.title(),
                heading_style
            )
        )

        for paragraph_text in content.splitlines():

            paragraph_text = paragraph_text.strip()

            if not paragraph_text:
                continue

            paragraph_text = html.escape(
                paragraph_text
            )

            if paragraph_text.startswith("-"):

                paragraph_text = (
                    "• "
                    + paragraph_text[1:].strip()
                )

            story.append(
                Paragraph(
                    paragraph_text,
                    body_style
                )
            )

    sources = sections.get(
        "SOURCE RECORDS",
        ""
    )

    if sources:

        story.append(
            Paragraph(
                "Research Evidence & Sources",
                heading_style
            )
        )

        for line in sources.splitlines():

            line = line.strip()

            if not line:
                continue

            line = html.escape(line)

            story.append(
                Paragraph(
                    "• " + line,
                    body_style
                )
            )

    document.build(story)

    output.seek(0)

    return output.getvalue()


# ============================================================
# GLOBAL STATISTICS
# ============================================================

total_records = sum(
    len(df)
    for df in data.values()
)

available_collections = sum(
    1
    for df in data.values()
    if not df.empty
)

source_counts = {}

for filename, df in data.items():

    if df.empty:
        continue

    source = source_from_filename(filename)

    source_counts[source] = (
        source_counts.get(source, 0)
        + len(df)
    )


# ============================================================
# SESSION STATE
# ============================================================

if "search_query" not in st.session_state:
    st.session_state["search_query"] = ""

if "search_results" not in st.session_state:
    st.session_state["search_results"] = []

if "generated_article" not in st.session_state:
    st.session_state["generated_article"] = ""

if "generated_image" not in st.session_state:
    st.session_state["generated_image"] = None

if "generated_other" not in st.session_state:
    st.session_state["generated_other"] = ""

if "generated_query" not in st.session_state:
    st.session_state["generated_query"] = ""

if "polar_ai_messages" not in st.session_state:
    st.session_state["polar_ai_messages"] = []

if "polar_ai_question" not in st.session_state:
    st.session_state["polar_ai_question"] = ""


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.title("POLAR BHARAT")

    st.caption(
        "Indian Polar Research Knowledge and Outreach Portal"
    )

    st.divider()

    st.subheader("Portal")

    page = st.radio(
        "Navigation",
        [
            "Home",
            "Archive Explorer",
            "Research Search",
            "Research Overview",
            "Generate Content",
            "Station Explorer",
            "Polar AI",
            "User Manual"
        ],
        label_visibility="collapsed"
    )

    st.divider()

    st.subheader("Knowledge Sources")

    st.write("NCPOR")
    st.write("NPDC")
    st.write("Research Publications")
    st.write("Expedition Reports")
    st.write("Polar Images")

    st.divider()

    st.caption(
        "Connected to the Polar Research Knowledge Base"
    )


# ============================================================
# TOP HEADER
# ============================================================

st.title("POLAR BHARAT")

st.caption(
    "Indian Polar Research Knowledge and Outreach Portal"
)

st.divider()


# ============================================================
# HOME
# ============================================================

if page == "Home":

    st.header("Indian Polar Science")

    st.write(
        "Explore research, datasets, publications, expedition "
        "records, stations and scientific knowledge from India's "
        "polar research ecosystem."
    )

    st.divider()

    st.subheader("Research Search")

    home_query = st.text_input(
        "Search the polar research archive",
        value=st.session_state.get(
            "search_query",
            ""
        ),
        placeholder=(
            "Example: Arctic climate research"
        ),
        key="home_search"
    )

    if st.button(
        "Search Research",
        type="primary",
        key="home_search_button"
    ):

        if home_query.strip():

            st.session_state["search_query"] = (
                home_query.strip()
            )

            with st.spinner(
                "Searching the research knowledge base..."
            ):

                results = search_records(
                    home_query.strip()
                )

            st.session_state["search_results"] = results

            if results:

                st.success(
                    f"Found {len(results)} relevant research records."
                )

                st.info(
                    "Open Research Search from the sidebar to "
                    "view the complete results."
                )

            else:

                st.warning(
                    "No relevant research records were found."
                )

        else:

            st.warning(
                "Please enter a research topic first."
            )

    st.divider()

    st.subheader("Main Topics")

    st.write(
        "Select a topic to quickly search the research archive."
    )

    topic_columns = st.columns(4)

    for index, topic in enumerate(MAIN_TOPICS):

        with topic_columns[index % 4]:

            if st.button(
                topic,
                key=f"home_topic_{index}",
                use_container_width=True
            ):

                st.session_state["search_query"] = topic

                with st.spinner(
                    f"Searching for {topic}..."
                ):

                    results = search_records(
                        topic
                    )

                st.session_state["search_results"] = results

                if results:

                    st.success(
                        f"{len(results)} records found for {topic}."
                    )

                else:

                    st.warning(
                        f"No records found for {topic}."
                    )

    st.divider()

    st.subheader("Quick Access")

    q1, q2, q3 = st.columns(3)

    with q1:

        with st.container(border=True):

            st.subheader("Research Search")

            st.write(
                "Search the knowledge base using natural-language questions."
            )

            if st.button(
                "Open Research Search",
                key="home_open_search",
                use_container_width=True
            ):

                st.session_state["search_query"] = ""

                st.info(
                    "Select Research Search from the sidebar."
                )

    with q2:

        with st.container(border=True):

            st.subheader("Scientific Datasets")

            st.write(
                "Explore datasets and scientific metadata collected "
                "from the available research sources."
            )

            if st.button(
                "Open Archive Explorer",
                key="home_open_archive",
                use_container_width=True
            ):

                st.info(
                    "Select Archive Explorer from the sidebar."
                )

    with q3:

        with st.container(border=True):

            st.subheader("Publications")

            st.write(
                "Browse publications and related research records."
            )

            if st.button(
                "Browse Publications",
                key="home_open_publications",
                use_container_width=True
            ):

                st.session_state["archive_collection"] = (
                    "polar_publications.csv"
                )

                st.info(
                    "Select Archive Explorer from the sidebar "
                    "to browse publications."
                )

    st.divider()

    st.subheader("Archive at a Glance")

    metric1, metric2, metric3, metric4 = st.columns(4)

    with metric1:

        st.metric(
            "Total Records",
            f"{total_records:,}"
        )

    with metric2:

        st.metric(
            "Collections",
            len(DATA_FILES)
        )

    with metric3:

        st.metric(
            "Loaded Collections",
            available_collections
        )

    with metric4:

        st.metric(
            "Research Sources",
            len(source_counts)
        )


# ============================================================
# ARCHIVE EXPLORER
# ============================================================

elif page == "Archive Explorer":

    st.header("Polar Research Archive")

    st.write(
        "Explore research records, datasets, publications, projects, "
        "station metadata, images and expedition resources."
    )

    st.divider()

    metric1, metric2, metric3, metric4 = st.columns(4)

    with metric1:

        st.metric(
            "Total Records",
            f"{total_records:,}"
        )

    with metric2:

        st.metric(
            "Collections",
            len(DATA_FILES)
        )

    with metric3:

        st.metric(
            "Loaded Collections",
            available_collections
        )

    with metric4:

        st.metric(
            "Research Sources",
            len(source_counts)
        )

    st.divider()

    st.subheader("Research Collections")

    for filename in DATA_FILES:

        df = data.get(
            filename,
            pd.DataFrame()
        )

        with st.container(border=True):

            col1, col2 = st.columns(
                [4, 1]
            )

            with col1:

                st.subheader(
                    COLLECTION_NAMES.get(
                        filename,
                        filename
                    )
                )

                st.write(
                    COLLECTION_DESCRIPTIONS.get(
                        filename,
                        ""
                    )
                )

            with col2:

                st.metric(
                    "Records",
                    f"{len(df):,}"
                )

                st.caption(
                    source_from_filename(filename)
                )

    st.divider()

    st.subheader("Collection Explorer")

    default_collection = st.session_state.get(
        "archive_collection",
        DATA_FILES[0]
    )

    if default_collection not in DATA_FILES:
        default_collection = DATA_FILES[0]

    selected_collection = st.selectbox(
        "Select a collection",
        DATA_FILES,
        index=DATA_FILES.index(
            default_collection
        ),
        format_func=lambda x:
            COLLECTION_NAMES.get(x, x)
    )

    st.session_state["archive_collection"] = (
        selected_collection
    )

    selected_df = data.get(
        selected_collection,
        pd.DataFrame()
    )

    if selected_df.empty:

        st.warning(
            "This collection is empty or could not be loaded."
        )

    else:

        c1, c2, c3 = st.columns(3)

        with c1:

            st.metric(
                "Records",
                f"{len(selected_df):,}"
            )

        with c2:

            st.metric(
                "Fields",
                len(selected_df.columns)
            )

        with c3:

            st.metric(
                "Source",
                source_from_filename(
                    selected_collection
                )
            )

        st.divider()

        st.subheader("Filter Collection")

        search_text = st.text_input(
            "Search records",
            placeholder=(
                "Search across all available fields..."
            ),
            key="archive_search"
        )

        rows_to_show = st.selectbox(
            "Rows to display",
            [
                10,
                25,
                50,
                100,
                250,
                500
            ],
            index=2
        )

        filtered_df = selected_df.copy()

        if search_text.strip():

            search_value = (
                search_text
                .lower()
                .strip()
            )

            mask = filtered_df.apply(
                lambda row:
                    row.astype(str)
                    .str.lower()
                    .str.contains(
                        search_value,
                        regex=False
                    )
                    .any(),
                axis=1
            )

            filtered_df = filtered_df[
                mask
            ]

        filterable_columns = []

        for column in selected_df.columns:

            try:

                unique_count = selected_df[
                    column
                ].nunique(
                    dropna=True
                )

            except Exception:

                unique_count = 0

            if (
                unique_count > 1
                and unique_count <= 30
            ):

                filterable_columns.append(
                    column
                )

        if filterable_columns:

            st.subheader(
                "Additional Filters"
            )

            number_of_filters = min(
                len(filterable_columns),
                3
            )

            filter_columns = st.columns(
                number_of_filters
            )

            for i, column in enumerate(
                filterable_columns[:3]
            ):

                values = sorted(
                    [
                        str(value)
                        for value in selected_df[
                            column
                        ].dropna().unique()
                    ]
                )

                with filter_columns[i]:

                    selected_value = st.selectbox(
                        column,
                        ["All"] + values,
                        key=(
                            "archive_filter_"
                            + str(selected_collection)
                            + "_"
                            + str(column)
                        )
                    )

                    if selected_value != "All":

                        filtered_df = filtered_df[
                            filtered_df[column]
                            .astype(str)
                            == selected_value
                        ]

        st.divider()

        result_col1, result_col2 = st.columns(2)

        with result_col1:

            st.write(
                f"Matching records: **{len(filtered_df):,}**"
            )

        with result_col2:

            st.write(
                f"Displaying up to **{rows_to_show:,}** rows"
            )

        st.dataframe(
            filtered_df.head(rows_to_show),
            use_container_width=True,
            hide_index=True,
            height=550
        )


# ============================================================
# RESEARCH SEARCH
# ============================================================

elif page == "Research Search":

    st.header("Research Search")

    st.write(
        "Search across the polar research knowledge base "
        "using natural-language questions."
    )

    st.divider()

    default_query = st.session_state.get(
        "search_query",
        ""
    )

    query = st.text_input(
        "Research question",
        value=default_query,
        placeholder=(
            "Example: What datasets are available for Arctic climate research?"
        ),
        key="research_query_input"
    )

    if st.button(
        "Search Research",
        type="primary",
        key="research_search_button"
    ):

        if not query.strip():

            st.warning(
                "Please enter a research topic or question."
            )

        else:

            with st.spinner(
                "Searching research records..."
            ):

                results = search_records(
                    query.strip()
                )

            st.session_state["search_results"] = results
            st.session_state["search_query"] = query.strip()

            if not results:

                st.warning(
                    "No relevant research records were found."
                )

            else:

                st.success(
                    f"Found {len(results)} relevant research records."
                )

    results = st.session_state.get(
        "search_results",
        []
    )

    if results:

        st.divider()

        sources = set()
        content_types = set()

        for score, index, row in results:

            sources.add(
                get_source(row)
            )

            content_type = clean_text(
                row.get(
                    "content_type",
                    ""
                )
            )

            if not content_type:
                content_type = "Research Record"

            content_types.add(
                content_type
            )

        st.subheader("Search Summary")

        col1, col2, col3 = st.columns(3)

        with col1:

            st.metric(
                "Relevant Records",
                len(results)
            )

        with col2:

            st.metric(
                "Sources",
                len(sources)
            )

        with col3:

            st.metric(
                "Content Types",
                len(content_types)
            )

        st.divider()

        st.subheader("Research Results")

        for number, (
            score,
            index,
            row
        ) in enumerate(
            results,
            start=1
        ):

            title = get_title(row)

            if not title:

                title = "Untitled Research Record"

            source = get_source(row)

            content_type = clean_text(
                row.get(
                    "content_type",
                    ""
                )
            )

            if not content_type:

                content_type = "Research Record"

            category = clean_text(
                row.get(
                    "category",
                    ""
                )
            )

            research_area = clean_text(
                row.get(
                    "research_area",
                    ""
                )
            )

            topic = clean_text(
                row.get(
                    "topic",
                    ""
                )
            )

            if not topic:

                topic = clean_text(
                    row.get(
                        "subcategory",
                        ""
                    )
                )

            station = clean_text(
                row.get(
                    "station",
                    ""
                )
            )

            description = clean_text(
                row.get(
                    "description",
                    ""
                )
            )

            url = clean_text(
                row.get(
                    "metadata_url",
                    ""
                )
            )

            if not url:

                url = clean_text(
                    row.get(
                        "url",
                        ""
                    )
                )

            with st.container(border=True):

                st.caption(
                    f"RESULT {number:02d}"
                )

                st.subheader(title)

                info_col1, info_col2 = st.columns(2)

                with info_col1:

                    st.write(
                        f"**Source:** {source}"
                    )

                    st.write(
                        f"**Type:** {content_type}"
                    )

                    if category:

                        st.write(
                            f"**Category:** {category}"
                        )

                    if research_area:

                        st.write(
                            f"**Research Area:** {research_area}"
                        )

                with info_col2:

                    if topic:

                        st.write(
                            f"**Topic:** {topic}"
                        )

                    if station:

                        st.write(
                            f"**Station:** {station}"
                        )

                    st.write(
                        f"**Search Score:** {score}"
                    )

                if description:

                    with st.expander(
                        "View description"
                    ):

                        st.write(
                            description
                        )

                if (
                    url
                    and url.startswith("http")
                ):

                    st.link_button(
                        "View Source",
                        url
                    )

        st.divider()

        st.subheader("Research Analysis")

        analysis = analyze_records(
            results,
            st.session_state.get(
                "search_query",
                ""
            )
        )

        if analysis:

            col1, col2 = st.columns(2)

            with col1:

                st.markdown("#### Sources")

                for source, count in (
                    analysis["sources"].items()
                ):

                    st.write(
                        f"**{source}** — {count} records"
                    )

            with col2:

                st.markdown("#### Content Types")

                for content_type, count in (
                    analysis["content_types"].items()
                ):

                    st.write(
                        f"**{content_type}** — {count} records"
                    )

            st.markdown("#### Research Areas")

            if analysis["categories"]:

                for category, count in (
                    analysis["categories"][:10]
                ):

                    st.write(
                        f"**{category}** — {count} records"
                    )

            else:

                st.write(
                    "No research areas identified."
                )

            st.markdown("#### Geographic Focus")

            if analysis["locations"]:

                location_columns = st.columns(
                    min(
                        len(analysis["locations"]),
                        4
                    )
                )

                for i, location in enumerate(
                    analysis["locations"]
                ):

                    location_columns[
                        i % len(location_columns)
                    ].info(
                        location
                    )

            else:

                st.write(
                    "No clear geographic focus detected."
                )


# ============================================================
# GENERATE CONTENT
# ============================================================

elif page == "Generate Content":

    st.header("Generate Outreach Content")

    st.write(
        "Transform retrieved polar research into articles and "
        "outreach content using the available research records."
    )

    st.info(
        "AI-generated content should be verified against the "
        "original research records before publication."
    )

    st.divider()

    query = st.text_input(
        "Research topic",
        value=st.session_state.get(
            "search_query",
            ""
        ),
        placeholder=(
            "Example: Indian research in Antarctica"
        ),
        key="generate_query"
    )

    content_type = st.selectbox(
        "Content format",
        [
            "Website Article",
            "Social Media Post",
            "Short Outreach Post",
            "All Formats"
        ]
    )

    content_map = {

        "Website Article":
            ["article"],

        "Social Media Post":
            ["social"],

        "Short Outreach Post":
            ["short"],

        "All Formats":
            [
                "article",
                "social",
                "short"
            ]
    }

    requested_content = content_map[
        content_type
    ]

    stored_results = st.session_state.get(
        "search_results",
        []
    )

    if stored_results:

        st.success(
            f"{len(stored_results)} research records are "
            "available from your latest search."
        )

        use_existing = st.checkbox(
            "Use results from my latest search",
            value=True
        )

    else:

        use_existing = False

        st.info(
            "No previous search results are available. "
            "The portal will search using the topic you enter."
        )

    st.divider()

    if st.button(
        "Generate Content",
        type="primary",
        key="generate_content_button"
    ):

        if not query.strip():

            st.warning(
                "Please enter a research topic."
            )

        else:

            if (
                use_existing
                and stored_results
            ):

                results = stored_results

            else:

                with st.spinner(
                    "Searching research records..."
                ):

                    results = search_records(
                        query.strip()
                    )

            if not results:

                st.warning(
                    "No relevant research records were found."
                )

            else:

                st.session_state[
                    "generated_query"
                ] = query.strip()

                st.session_state[
                    "generated_results"
                ] = results

                if "article" in requested_content:

                    with st.spinner(
                        "Writing evidence-based article..."
                    ):

                        article = generate_article(
                            query.strip(),
                            results
                        )

                    if article:

                        st.session_state[
                            "generated_article"
                        ] = article

                        with st.spinner(
                            "Creating research illustration..."
                        ):

                            image_bytes = (
                                generate_research_image(
                                    query.strip(),
                                    results
                                )
                            )

                        st.session_state[
                            "generated_image"
                        ] = image_bytes

                    else:

                        st.error(
                            "Article generation failed."
                        )

                if (
                    "social" in requested_content
                    or "short" in requested_content
                ):

                    with st.spinner(
                        "Generating outreach content..."
                    ):

                        generated = generate_outreach(
                            query.strip(),
                            results,
                            requested_content
                        )

                    st.session_state[
                        "generated_other"
                    ] = generated

                st.success(
                    "Content generation completed."
                )

    article = st.session_state.get(
        "generated_article",
        ""
    )

    image_bytes = st.session_state.get(
        "generated_image",
        None
    )

    if article:

        st.divider()

        st.subheader("Article Preview")

        article_sections = parse_article(
            article
        )

        if image_bytes:

            st.image(
                image_bytes,
                use_container_width=True
            )

        title = article_sections.get(
            "TITLE",
            "Polar Research Article"
        )

        st.title(title)

        subtitle = article_sections.get(
            "SUBTITLE",
            ""
        )

        if subtitle:

            st.write(
                f"_{subtitle}_"
            )

        display_sections = [
            "INTRODUCTION",
            "RESEARCH HIGHLIGHTS",
            "MAIN STORY",
            "GEOGRAPHIC CONTEXT",
            "WHY IT MATTERS",
            "CONCLUSION"
        ]

        for section_name in display_sections:

            content = article_sections.get(
                section_name,
                ""
            )

            if not content:
                continue

            st.subheader(
                section_name.title()
            )

            st.markdown(
                content
            )

        sources = article_sections.get(
            "SOURCE RECORDS",
            ""
        )

        if sources:

            st.divider()

            st.subheader(
                "Research Evidence and Sources"
            )

            st.markdown(
                sources
            )

        st.divider()

        st.subheader("Export Article")

        docx_bytes = create_docx(
            article,
            image_bytes
        )

        pdf_bytes = create_pdf(
            article,
            image_bytes
        )

        col1, col2 = st.columns(2)

        with col1:

            st.download_button(
                "Download DOCX",
                data=docx_bytes,
                file_name="polar_research_article.docx",
                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml.document"
                ),
                use_container_width=True
            )

        with col2:

            st.download_button(
                "Download PDF",
                data=pdf_bytes,
                file_name="polar_research_article.pdf",
                mime="application/pdf",
                use_container_width=True
            )

    other_content = st.session_state.get(
        "generated_other",
        ""
    )

    if other_content:

        st.divider()

        st.subheader(
            "Additional Outreach Content"
        )

        with st.container(border=True):

            st.markdown(
                other_content
            )


# ============================================================
# RESEARCH OVERVIEW
# ============================================================

elif page == "Research Overview":

    st.header("Research Overview")

    st.write(
        "A high-level view of the research information "
        "currently available across the loaded collections."
    )

    st.divider()

    col1, col2, col3, col4 = st.columns(4)

    with col1:

        st.metric(
            "Total Records",
            f"{total_records:,}"
        )

    with col2:

        st.metric(
            "Collections",
            len(DATA_FILES)
        )

    with col3:

        st.metric(
            "Loaded Collections",
            available_collections
        )

    with col4:

        st.metric(
            "Research Sources",
            len(source_counts)
        )

    st.divider()

    st.subheader("Collection Distribution")

    distribution_rows = []

    for filename in DATA_FILES:

        df = data.get(
            filename,
            pd.DataFrame()
        )

        distribution_rows.append(
            {
                "Collection":
                    COLLECTION_NAMES.get(
                        filename,
                        filename
                    ),

                "Records":
                    len(df)
            }
        )

    distribution_df = pd.DataFrame(
        distribution_rows
    )

    st.bar_chart(
        distribution_df.set_index(
            "Collection"
        ),
        height=400
    )

    st.subheader("Research Source Distribution")

    source_df = pd.DataFrame(
        {
            "Source":
                list(
                    source_counts.keys()
                ),

            "Records":
                list(
                    source_counts.values()
                )
        }
    )

    if not source_df.empty:

        st.bar_chart(
            source_df.set_index(
                "Source"
            ),
            height=300
        )

        st.dataframe(
            source_df,
            use_container_width=True,
            hide_index=True
        )

    st.subheader("Content Type Distribution")

    content_type_counts = {}

    for filename, df in data.items():

        if df.empty:
            continue

        if "content_type" not in df.columns:
            continue

        for value in df[
            "content_type"
        ].dropna():

            value = str(
                value
            ).strip()

            if not value:
                continue

            content_type_counts[value] = (
                content_type_counts.get(
                    value,
                    0
                )
                + 1
            )

    if content_type_counts:

        content_type_df = pd.DataFrame(
            {
                "Content Type":
                    list(
                        content_type_counts.keys()
                    ),

                "Records":
                    list(
                        content_type_counts.values()
                    )
            }
        )

        st.bar_chart(
            content_type_df.set_index(
                "Content Type"
            ),
            height=350
        )

    else:

        st.info(
            "Content type information is not available "
            "in the loaded collections."
        )

    st.subheader("Collection Details")

    details = []

    for filename in DATA_FILES:

        df = data.get(
            filename,
            pd.DataFrame()
        )

        details.append(
            {
                "Collection":
                    COLLECTION_NAMES.get(
                        filename,
                        filename
                    ),

                "Source":
                    source_from_filename(
                        filename
                    ),

                "Records":
                    len(df),

                "Fields":
                    len(df.columns)
                    if not df.empty
                    else 0,

                "Status":
                    "Loaded"
                    if not df.empty
                    else "Unavailable"
            }
        )

    details_df = pd.DataFrame(
        details
    )

    st.dataframe(
        details_df,
        use_container_width=True,
        hide_index=True
    )

    st.subheader("Archive Insights")

    if not distribution_df.empty:

        largest_collection = (
            distribution_df
            .sort_values(
                "Records",
                ascending=False
            )
            .iloc[0]
        )

        st.info(
            f"The largest collection is "
            f"**{largest_collection['Collection']}**, "
            f"with **{int(largest_collection['Records']):,} "
            f"records."
        )

        if source_counts:

            largest_source = max(
                source_counts,
                key=source_counts.get
            )

            st.info(
                f"The source with the most archived records is "
                f"**{largest_source}**, with "
                f"**{source_counts[largest_source]:,} records."
            )


# ============================================================
# STATION EXPLORER
# ============================================================

elif page == "Station Explorer":

    st.header("Station Explorer")

    st.write(
        "Explore station metadata available in the current "
        "polar research archive."
    )

    st.divider()

    station_file = "npdc_station_metadata.csv"

    station_df = data.get(
        station_file,
        pd.DataFrame()
    )

    if station_df.empty:

        st.warning(
            "Station metadata is currently unavailable."
        )

    else:

        st.success(
            f"{len(station_df):,} station metadata records loaded."
        )

        st.subheader("Station Search")

        station_search = st.text_input(
            "Search stations",
            placeholder=(
                "Search by station name, location or metadata..."
            )
        )

        display_station_df = station_df.copy()

        if station_search.strip():

            search_value = (
                station_search
                .lower()
                .strip()
            )

            mask = display_station_df.apply(
                lambda row:
                    row.astype(str)
                    .str.lower()
                    .str.contains(
                        search_value,
                        regex=False
                    )
                    .any(),
                axis=1
            )

            display_station_df = (
                display_station_df[mask]
            )

        st.dataframe(
            display_station_df,
            use_container_width=True,
            hide_index=True,
            height=550
        )


# ============================================================
# POLAR AI
# ============================================================

elif page == "Polar AI":

    st.header("Polar AI")

    st.write(
        "A research-oriented assistant connected to the "
        "Polar Research Knowledge Base."
    )

    st.info(
        "Polar AI retrieves relevant research records first. "
        "Use the underlying records and original sources when "
        "verifying scientific information."
    )

    st.divider()

    st.subheader("Suggested Questions")

    suggestions = [
        "Show me Arctic datasets",
        "What research projects are available?",
        "Summarize research related to Maitri Station",
        "Tell me about atmospheric research in Antarctica"
    ]

    suggestion_columns = st.columns(2)

    for i, suggestion in enumerate(suggestions):

        with suggestion_columns[i % 2]:

            if st.button(
                suggestion,
                key=f"polar_suggestion_{i}",
                use_container_width=True
            ):

                st.session_state[
                    "polar_ai_question"
                ] = suggestion

                st.rerun()

    question = st.text_input(
        "Ask Polar AI",
        value=st.session_state.get(
            "polar_ai_question",
            ""
        ),
        placeholder=(
            "Ask about projects, datasets, publications, "
            "stations or research topics..."
        ),
        key="polar_ai_input"
    )

    ai_col1, ai_col2 = st.columns(
        [4, 1]
    )

    with ai_col1:

        ask_button = st.button(
            "Ask Polar AI",
            type="primary",
            use_container_width=True
        )

    with ai_col2:

        clear_button = st.button(
            "Clear Conversation",
            use_container_width=True
        )

    if clear_button:

        st.session_state[
            "polar_ai_messages"
        ] = []

        st.session_state[
            "polar_ai_question"
        ] = ""

        st.rerun()

    if ask_button:

        if not question.strip():

            st.warning(
                "Enter a question first."
            )

        else:

            with st.spinner(
                "Searching the Polar Research Knowledge Base..."
            ):

                results = search_records(
                    question.strip()
                )

            if not results:

                answer = (
                    "No relevant research records were found "
                    "for this question."
                )

            else:

                answer_lines = [
                    "I found relevant records in the "
                    "Polar Research Knowledge Base.",
                    "",
                    f"Retrieved records: {len(results)}",
                    ""
                ]

                for number, (
                    score,
                    index,
                    row
                ) in enumerate(
                    results[:5],
                    start=1
                ):

                    title = get_title(row)

                    if not title:

                        title = (
                            "Untitled Research Record"
                        )

                    source = get_source(row)

                    answer_lines.append(
                        f"{number}. {title} ({source})"
                    )

                answer = "\n".join(
                    answer_lines
                )

            st.session_state[
                "polar_ai_messages"
            ].append(
                {
                    "question":
                        question.strip(),

                    "answer":
                        answer
                }
            )

            st.session_state[
                "polar_ai_question"
            ] = ""

            st.rerun()

    if st.session_state[
        "polar_ai_messages"
    ]:

        st.divider()

        st.subheader("Conversation")

        for message in reversed(
            st.session_state[
                "polar_ai_messages"
            ]
        ):

            with st.container(border=True):

                st.markdown(
                    "#### You"
                )

                st.write(
                    message["question"]
                )

                st.divider()

                st.markdown(
                    "#### Polar AI"
                )

                st.write(
                    message["answer"]
                )

                st.caption(
                    "Based on retrieved research records. "
                    "Verify critical scientific information."
                )


# ============================================================
# USER MANUAL
# ============================================================

elif page == "User Manual":

    st.header("User Manual")

    st.write(
        "Use this guide to discover, inspect and generate "
        "content from the Polar Bharat research archive."
    )

    st.divider()

    manual_steps = [

        (
            "1",
            "Search Research",
            "Enter a natural-language question to find relevant "
            "projects, datasets, publications, stations and research records."
        ),

        (
            "2",
            "Browse the Archive",
            "Use Archive Explorer to inspect the collections currently "
            "loaded into the portal."
        ),

        (
            "3",
            "Filter Records",
            "Use text search and available categorical filters to "
            "narrow down a collection."
        ),

        (
            "4",
            "Inspect Sources",
            "Open individual research records and use View Source "
            "when an original source link is available."
        ),

        (
            "5",
            "Understand the Archive",
            "Use Research Overview to see collection sizes, source "
            "distribution and content-type information."
        ),

        (
            "6",
            "Explore Stations",
            "Use Station Explorer to search and inspect available "
            "polar station metadata."
        ),

        (
            "7",
            "Use Polar AI",
            "Ask research-oriented questions and inspect the "
            "records retrieved for the answer."
        ),

        (
            "8",
            "Generate Outreach",
            "Use Generate Content to turn retrieved research into "
            "articles and social or outreach content."
        ),

        (
            "9",
            "Export Articles",
            "Generated articles can be exported as DOCX or PDF files."
        ),

        (
            "10",
            "Verify Information",
            "Always verify important scientific claims against the "
            "underlying research records and original sources."
        )
    ]

    for number, title, description in manual_steps:

        with st.container(border=True):

            col1, col2 = st.columns(
                [1, 8]
            )

            with col1:

                st.subheader(
                    number
                )

            with col2:

                st.subheader(
                    title
                )

                st.write(
                    description
                )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "POLAR BHARAT | Indian Polar Research Knowledge and Outreach Portal"
)

st.caption(
    "Connecting people with India's polar science."
)

st.caption(
    "NCPOR + NPDC Knowledge Ecosystem | 2026"
)